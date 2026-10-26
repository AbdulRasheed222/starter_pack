"""Turn uploaded files into plain-text sections with page and heading metadata.

PDFs are read with pypdf (BSD licensed). PyMuPDF extracts layout better but is AGPL, which
would bind every product built on this starter.
"""

import io
import re
from dataclasses import dataclass
from pathlib import PurePath

import docx
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class ParseError(ValueError):
    """The file could not be turned into text. The message is safe to show to the user."""


@dataclass
class Section:
    text: str
    page: int | None = None
    heading: str | None = None


@dataclass
class ParsedDocument:
    title: str
    kind: str
    sections: list[Section]
    pages: int


KINDS_BY_SUFFIX = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".md": "markdown",
    ".markdown": "markdown",
    ".txt": "text",
}

_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def kind_for_filename(filename: str) -> str:
    kind = KINDS_BY_SUFFIX.get(PurePath(filename).suffix.lower())
    if kind is None:
        supported = ", ".join(sorted(KINDS_BY_SUFFIX))
        raise ParseError(f"Unsupported file type. Supported types: {supported}")
    return kind


def parse_file(filename: str, data: bytes) -> ParsedDocument:
    kind = kind_for_filename(filename)
    title = PurePath(filename).stem
    if kind == "pdf":
        return _parse_pdf(title, data)
    if kind == "docx":
        return _parse_docx(title, data)
    text = _decode(data)
    if kind == "markdown":
        return ParsedDocument(title, kind, parse_markdown(text), pages=1)
    return ParsedDocument(title, kind, [Section(text.strip())] if text.strip() else [], pages=1)


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    # cp1254 is the Windows Turkish code page, still common in older text files.
    for encoding in ("utf-8-sig", "cp1254"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _parse_pdf(title: str, data: bytes) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ParseError("This PDF is password protected.")
        sections = [
            Section(text=(page.extract_text() or "").strip(), page=number)
            for number, page in enumerate(reader.pages, start=1)
        ]
        pages = len(reader.pages)
        meta_title = (reader.metadata.title if reader.metadata else None) or ""
    except ParseError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise ParseError("This PDF could not be opened.") from exc
    sections = [s for s in sections if s.text]
    if not sections:
        raise ParseError(
            "This PDF has no text layer, probably because it is scanned. "
            "Scanned PDFs are not supported."
        )
    return ParsedDocument(str(meta_title).strip() or title, "pdf", sections, pages)


def _parse_docx(title: str, data: bytes) -> ParsedDocument:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for bad files
        raise ParseError("This Word file could not be opened.") from exc
    sections: list[Section] = []
    heading: str | None = None
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style is not None and paragraph.style.name.lower().startswith("heading"):
            if lines:
                sections.append(Section("\n".join(lines), heading=heading))
            heading, lines = text, []
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    if lines:
        sections.append(Section("\n".join(lines), heading=heading))
    return ParsedDocument(title, "docx", sections, pages=1)


def parse_markdown(text: str) -> list[Section]:
    """One section per heading, keeping the nearest heading as context."""
    sections: list[Section] = []
    heading: str | None = None
    lines: list[str] = []
    for line in text.splitlines():
        match = _HEADING.match(line)
        if match:
            if "".join(lines).strip():
                sections.append(Section("\n".join(lines).strip(), heading=heading))
            heading, lines = match.group(2), []
        else:
            lines.append(line)
    if "".join(lines).strip():
        sections.append(Section("\n".join(lines).strip(), heading=heading))
    return sections
